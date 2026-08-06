# -*- coding: utf-8 -*-
# Gera o template da planilha do dash de Eleicoes. Time da Globo preenche; front espelha.
# Cada aba = uma secao do dash. Nao renomear cabecalhos (o front le por eles).
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

SAIDA = "Planilha_Dash_Eleicoes_TEMPLATE.xlsx"

# aba -> (colunas, linhas de exemplo)
ABAS = {
    "TOTALS": (
        ["data", "posts", "interacoes", "views", "alcance"],
        [["06/08", 120, 450000, 3200000, 2100000],
         ["07/08", 138, 512000, 3600000, 2400000]],
    ),
    "PLATAFORMAS": (
        ["data", "plataforma", "posts", "interacoes", "views", "alcance"],
        [["06/08", "Instagram", 40, 210000, 1500000, 1100000],
         ["06/08", "TikTok", 25, 120000, 900000, ""],
         ["06/08", "X/Twitter", 35, 60000, 400000, ""],
         ["06/08", "YouTube", 12, 40000, 300000, ""],
         ["06/08", "Facebook", 8, 20000, 100000, 80000]],
    ),
    "TOP_POSTS": (
        ["data", "rank", "plataforma", "horario", "interacoes", "views", "alcance", "texto", "url"],
        [["06/08", 1, "Instagram", "20h15", 85000, 1200000, 900000,
          "Sabatina com o candidato X na GloboNews", "https://instagram.com/p/xxxx"],
         ["06/08", 2, "TikTok", "21h40", 62000, 800000, "",
          "Melhor momento do GloboNews Debate", "https://tiktok.com/@globonews/video/xxx"]],
    ),
    "DIARIO": (
        ["data", "mencoes", "contas", "engajamento", "sentimento_pos", "sentimento_neg", "tts"],
        [["06/08", 48200, 31000, 210000, "42%", "18%", 3],
         ["07/08", 61500, 39000, 265000, "39%", "22%", 5]],
    ),
    "SOV": (
        ["data", "ator", "percentual"],
        [["06/08", "GloboNews", 46],
         ["06/08", "CNN Brasil", 28],
         ["06/08", "Record News", 15],
         ["06/08", "Outros", 11]],
    ),
    "MENCOES_HORA": (
        ["data", "hora", "volume", "nota"],
        [["06/08", "18:00", 2100, ""],
         ["06/08", "20:00", 8600, "Pico durante a sabatina"],
         ["06/08", "22:00", 4300, ""]],
    ),
    "NUVEM_TERMOS": (
        ["data", "escopo", "termo", "peso"],
        [["06/08", "geral", "sabatina", 120],
         ["06/08", "geral", "economia", 90],
         ["06/08", "jornalista:Andreia Sadi", "firme", 40],
         ["06/08", "programa:GloboNews Debate", "bem organizado", 35]],
    ),
    "ASSOCIACOES": (
        ["data", "pauta", "marca", "mencoes"],
        [["06/08", "Debate", "g1", 1800],
         ["06/08", "Entrevistas", "GloboNews", 2400]],
    ),
    "JORNALISTAS": (
        ["data", "jornalista", "mencoes", "sentimento", "nota"],
        [["06/08", "Andreia Sadi", 5200, "positivo", "Elogiada pela condução da sabatina"],
         ["06/08", "Natuza Nery", 4100, "positivo", "Destaque nas perguntas sobre economia"]],
    ),
    "PROGRAMAS": (
        ["data", "programa", "mencoes", "nota"],
        [["06/08", "GloboNews Debate", 12800, "Maior volume do dia"],
         ["06/08", "Central das Eleicoes", 3400, "Estreia prevista 2a quinzena de agosto"]],
    ),
    "DIAGNOSTICO": (
        ["data", "eixo", "tipo", "titulo", "descricao"],
        [["06/08", "globonews", "funcionou", "Cobertura ao vivo da sabatina",
          "Transmissao simultanea puxou pico de mencoes as 20h"],
         ["06/08", "globonews", "nao_funcionou", "Corte tardio para os melhores momentos",
          "Clipes so sairam no dia seguinte, perderam a janela de busca"],
         ["06/08", "jornalistas", "funcionou", "Conducao da Andreia Sadi",
          "Perguntas diretas geraram recortes espontaneos positivos"]],
    ),
    "PONTOS": (
        ["data", "tipo", "titulo", "descricao"],
        [["06/08", "positivo", "Boa repercussao da sabatina",
          "Sentimento majoritariamente positivo sobre a conducao"],
         ["06/08", "atencao", "Criticas sobre tempo de fala",
          "Reclamacoes de desequilibrio no tempo entre candidatos, monitorar"]],
    ),
    "APRENDIZADOS": (
        ["data", "eixo", "texto"],
        [["06/08", "globonews", "Publicar os melhores momentos ainda durante o programa segura o pico de busca."],
         ["06/08", "jornalistas", "Perguntas objetivas dos ancoras geram mais recorte espontaneo que blocos longos."]],
    ),
}

INSTRUCOES = [
    ["Dash Eleicoes GloboNews - Template de preenchimento"],
    [""],
    ["Como usar:"],
    ["1. Cada aba abaixo alimenta uma secao do dashboard. NAO renomeie os cabecalhos (linha 1)."],
    ["2. A coluna 'data' usa o formato DD/MM (ex: 06/08). Uma data por linha/registro."],
    ["3. As linhas de exemplo (a partir da linha 2) devem ser APAGADAS antes de usar de verdade."],
    ["4. Numeros sem pontuacao de milhar e sem unidade: 3200000, nao 3,2M."],
    ["5. Campos que voce nao tiver no dia, deixe em branco (o dash trata)."],
    [""],
    ["Abas e para que servem:"],
    ["TOTALS ........ Big Numbers de midia propria (posts, interacoes, views, alcance) por dia"],
    ["PLATAFORMAS ... quebra da midia propria por rede social"],
    ["TOP_POSTS ..... Top 10 posts proprios do dia, ranqueados por interacoes"],
    ["DIARIO ........ Big Numbers de midia espontanea (mencoes, contas, engajamento, sentimento, TTs)"],
    ["SOV ........... Share of Voice: participacao de cada ator no volume (defina os atores)"],
    ["MENCOES_HORA .. volume de mencoes por hora; 'nota' destaca o pico"],
    ["NUVEM_TERMOS .. termos + peso. 'escopo' = geral, jornalista:Nome ou programa:Nome"],
    ["ASSOCIACOES ... associacao pauta x marca (ex: Debate + g1)"],
    ["JORNALISTAS ... volume de mencoes por jornalista/apresentador + sentimento"],
    ["PROGRAMAS ..... repercussao dos programas (GloboNews Debate, Central das Eleicoes)"],
    ["DIAGNOSTICO ... 'eixo' = globonews ou jornalistas; 'tipo' = funcionou ou nao_funcionou"],
    ["PONTOS ........ 'tipo' = positivo ou atencao (crises/pontos de atencao entram como atencao)"],
    ["APRENDIZADOS .. aprendizado do dia; 'eixo' = globonews ou jornalistas"],
]

wb = openpyxl.Workbook()
HFILL = PatternFill("solid", fgColor="1B3A5B")
HFONT = Font(bold=True, color="FFFFFF")
EXFONT = Font(italic=True, color="888888")

ws0 = wb.active
ws0.title = "INSTRUCOES"
for r in INSTRUCOES:
    ws0.append(r)
ws0["A1"].font = Font(bold=True, size=14, color="1B3A5B")
ws0.column_dimensions["A"].width = 95

for aba, (cols, exemplos) in ABAS.items():
    ws = wb.create_sheet(aba)
    ws.append(cols)
    for c in range(1, len(cols)+1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HFILL; cell.font = HFONT
        cell.alignment = Alignment(horizontal="left")
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = max(12, len(cols[c-1])+4)
    for ex in exemplos:
        ws.append(ex)
        for c in range(1, len(cols)+1):
            ws.cell(row=ws.max_row, column=c).font = EXFONT
    ws.freeze_panes = "A2"

wb.save(SAIDA)
print("OK ->", SAIDA, "|", len(ABAS), "abas + INSTRUCOES")
